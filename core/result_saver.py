"""
결과 저장 (JSON, Excel, 링크 파일)
"""
import json
from pathlib import Path
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment


def save_content_to_json(site, result, save_folder=None, timestamp=None, log_func=None):
    """크롤링한 content를 JSON 파일로 저장"""
    def log(msg, level="info"):
        if log_func:
            log_func(msg, level)

    try:
        crawled_articles = result.get('crawled_articles', [])
        if not crawled_articles:
            log("⚠️ Content 저장 건너뜀: crawled_articles가 비어있음", "warning")
            return None

        site_name = site['name'].replace(' ', '_').replace('/', '_')

        base_dir = Path(save_folder) if save_folder else Path(".")
        output_dir = base_dir / "pagination_test_content" / site_name
        output_dir.mkdir(parents=True, exist_ok=True)

        if timestamp is None:
            timestamp = datetime.now().strftime("%Y%m%d-%H%M%S")
        json_file = output_dir / f"{timestamp}.json"

        failed_articles = result.get('failed_articles', [])

        content_data = {
            'site': {
                'name': site['name'],
                'url': site['url']
            },
            'test_info': {
                'timestamp': timestamp,
                'test_date': datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                'method_used': result.get('method_used', 'unknown'),
                'proxy_used': result.get('proxy_used', 'none'),
                'clicks': result.get('clicks', 0),
                'total_links': len(result.get('all_links', [])),
                'total_articles_crawled': len(crawled_articles),
                'total_articles_failed': len(failed_articles),
                'slider_detected': result.get('slider_detected', False),
                'slider_type': result.get('slider_type', '')
            },
            'articles': crawled_articles,
            'failed_articles': failed_articles
        }

        with open(json_file, 'w', encoding='utf-8') as f:
            json.dump(content_data, f, ensure_ascii=False, indent=2)

        log(f"\n💾 Content 저장 완료!", "success")
        log(f"   파일: {json_file}", "info")
        log(f"   게시글 수: {len(crawled_articles)}개", "info")
        if failed_articles:
            log(f"   ⚠️ 실패: {len(failed_articles)}개", "warning")

        return str(json_file)

    except Exception as e:
        log(f"⚠️ Content JSON 저장 실패: {e}", "warning")
        return None


def save_result_to_excel(site, result, save_folder=None, timestamp=None, log_func=None):
    """테스트 결과를 엑셀 파일로 저장 (카테고리 지원)"""
    def log(msg, level="info"):
        if log_func:
            log_func(msg, level)

    try:
        site_name = site['name'].replace(' ', '_').replace('/', '_')

        base_dir = Path(save_folder) if save_folder else Path(".")
        output_dir = base_dir / "pagination_test_results" / site_name
        output_dir.mkdir(parents=True, exist_ok=True)

        if timestamp is None:
            timestamp = datetime.now().strftime("%Y%m%d-%H%M%S")
        excel_file = output_dir / f"{timestamp}.xlsx"

        wb = Workbook()

        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")

        if 'category_results' in result:
            category_results = result.get('category_results', [])

            for idx, cat_result in enumerate(category_results):
                cat_name = cat_result.get('category_name', f'Category{idx+1}')
                cat_slug = cat_result.get('category_slug', f'cat{idx+1}')
                sheet_name = cat_slug[:31]

                ws = wb.active if idx == 0 else wb.create_sheet(title=sheet_name)
                if idx == 0:
                    ws.title = sheet_name

                headers = ["번호", "URL", "제목", "카테고리"]
                for col_idx, header in enumerate(headers, 1):
                    cell = ws.cell(row=1, column=col_idx, value=header)
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = Alignment(horizontal='center', vertical='center')

                cat_links = cat_result.get('all_links', [])
                for row_idx, link_data in enumerate(cat_links, 1):
                    url = link_data.get('url', '') if isinstance(link_data, dict) else str(link_data)
                    title = link_data.get('title', '') if isinstance(link_data, dict) else ''
                    ws.cell(row=row_idx+1, column=1, value=row_idx)
                    ws.cell(row=row_idx+1, column=2, value=url)
                    ws.cell(row=row_idx+1, column=3, value=title)
                    ws.cell(row=row_idx+1, column=4, value=cat_name)

                ws.column_dimensions['A'].width = 8
                ws.column_dimensions['B'].width = 80
                ws.column_dimensions['C'].width = 50
                ws.column_dimensions['D'].width = 30

            # 요약 시트
            ws_summary = wb.create_sheet(title="Summary")
            ws_summary.cell(row=1, column=1, value="카테고리별 요약").font = Font(bold=True, size=14)
            summary_headers = ["카테고리", "링크 수", "클릭 횟수"]
            for col_idx, header in enumerate(summary_headers, 1):
                cell = ws_summary.cell(row=3, column=col_idx, value=header)
                cell.fill = header_fill
                cell.font = header_font

            for row_idx, cat_result in enumerate(category_results, 1):
                ws_summary.cell(row=row_idx+3, column=1, value=cat_result.get('category_name', ''))
                ws_summary.cell(row=row_idx+3, column=2, value=len(cat_result.get('all_links', [])))
                ws_summary.cell(row=row_idx+3, column=3, value=cat_result.get('clicks', 0))

            total_row = len(category_results) + 5
            ws_summary.cell(row=total_row, column=1, value="전체").font = Font(bold=True)
            ws_summary.cell(row=total_row, column=2, value=result.get('final_count', 0)).font = Font(bold=True)
            ws_summary.cell(row=total_row, column=3, value=result.get('clicks', 0)).font = Font(bold=True)

        else:
            ws = wb.active
            ws.title = "URL List"

            headers = ["번호", "URL", "제목", "비고"]
            for col_idx, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col_idx, value=header)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal='center', vertical='center')

            all_links = result.get('all_links', [])
            for idx, link_data in enumerate(all_links, 1):
                url = link_data.get('url', '') if isinstance(link_data, dict) else str(link_data)
                title = link_data.get('title', '') if isinstance(link_data, dict) else ''
                ws.cell(row=idx+1, column=1, value=idx)
                ws.cell(row=idx+1, column=2, value=url)
                ws.cell(row=idx+1, column=3, value=title)

            ws.column_dimensions['A'].width = 8
            ws.column_dimensions['B'].width = 80
            ws.column_dimensions['C'].width = 50
            ws.column_dimensions['D'].width = 20

        # 메타데이터 시트
        ws_meta = wb.create_sheet(title="Test Info")
        meta_data = [
            ["테스트 정보", ""],
            ["사이트명", site['name']],
            ["URL", site['url']],
            ["테스트 일시", datetime.now().strftime("%Y-%m-%d %H:%M:%S")],
            ["", ""],
            ["결과 요약", ""],
        ]

        if 'category_results' in result:
            meta_data.extend([
                ["크롤링 방식", "카테고리별 페이지네이션"],
                ["총 카테고리 수", result.get('total_categories', 0)],
                ["성공 카테고리", result.get('success_categories', 0)],
                ["실패 카테고리", result.get('failed_categories', 0)],
            ])
        else:
            meta_data.extend([
                ["총 클릭 횟수", result.get('clicks', 0)],
                ["초기 게시글 수", result.get('initial_count', 0)],
                ["최종 게시글 수", result.get('final_count', 0)],
                ["증가한 게시글 수", result.get('new_articles', 0)],
            ])

        meta_data.extend([
            ["크롤링된 링크 수", len(result.get('all_links', []))],
            ["사용된 방법", result.get('method_used', 'unknown')],
        ])

        if result.get('slider_detected'):
            meta_data.extend([
                ["", ""],
                ["⚠️ 슬라이더/CAPTCHA 감지", ""],
                ["감지 여부", "예"],
                ["감지 타입", result.get('slider_type', 'unknown')],
            ])

        for row_idx, (key, value) in enumerate(meta_data, 1):
            ws_meta.cell(row=row_idx, column=1, value=key)
            ws_meta.cell(row=row_idx, column=2, value=value)
            if key in ["테스트 정보", "결과 요약"]:
                ws_meta.cell(row=row_idx, column=1).font = Font(bold=True, size=12)

        ws_meta.column_dimensions['A'].width = 20
        ws_meta.column_dimensions['B'].width = 60

        wb.save(excel_file)

        log(f"\n💾 결과 저장 완료!", "success")
        log(f"   파일: {excel_file}", "info")
        log(f"   링크 수: {len(result.get('all_links', []))}개", "info")

        return str(excel_file)

    except Exception as e:
        log(f"⚠️ 엑셀 저장 실패: {e}", "warning")
        return None


def save_links_to_file(site_name, all_links, config, save_folder=None, log_func=None):
    """링크 목록을 JSON 파일로 저장"""
    def log(msg, level="info"):
        if log_func:
            log_func(msg, level)

    try:
        safe_name = site_name.replace(' ', '_').replace('/', '_')
        base_dir = Path(save_folder) if save_folder else Path(".")
        output_dir = base_dir / "pagination_test_results" / safe_name
        output_dir.mkdir(parents=True, exist_ok=True)

        timestamp = datetime.now().strftime("%Y%m%d-%H%M%S")
        links_file = output_dir / f"{timestamp}_links.json"

        links_data = {
            'site_name': site_name,
            'timestamp': timestamp,
            'total_links': len(all_links),
            'config': config,
            'links': all_links
        }

        with open(links_file, 'w', encoding='utf-8') as f:
            json.dump(links_data, f, ensure_ascii=False, indent=2)

        log(f"📄 링크 파일 저장: {links_file} ({len(all_links)}개)", "success")
        return str(links_file)

    except Exception as e:
        log(f"⚠️ 링크 파일 저장 실패: {e}", "warning")
        return None


def save_phase2_result(crawled_articles, failed_articles, timestamp,
                       site_name="", save_folder=None, log_func=None):
    """Phase 2 크롤링 결과 별도 저장"""
    def log(msg, level="info"):
        if log_func:
            log_func(msg, level)

    try:
        safe_name = site_name.replace(' ', '_').replace('/', '_')
        base_dir = Path(save_folder) if save_folder else Path(".")
        output_dir = base_dir / "pagination_test_content" / safe_name
        output_dir.mkdir(parents=True, exist_ok=True)

        result_file = output_dir / f"{timestamp}_phase2.json"

        result_data = {
            'timestamp': timestamp,
            'total_crawled': len(crawled_articles),
            'total_failed': len(failed_articles),
            'articles': crawled_articles,
            'failed': failed_articles
        }

        with open(result_file, 'w', encoding='utf-8') as f:
            json.dump(result_data, f, ensure_ascii=False, indent=2)

        log(f"💾 Phase 2 결과 저장: {result_file}", "success")
        return str(result_file)

    except Exception as e:
        log(f"⚠️ Phase 2 결과 저장 실패: {e}", "warning")
        return None
