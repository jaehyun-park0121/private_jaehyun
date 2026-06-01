from __future__ import annotations

from math import ceil
from pathlib import Path
from typing import List, Optional, Sequence
import xml.etree.ElementTree as ET
import zipfile

from PyQt5.QtCore import Qt, QRect
from PyQt5.QtGui import QColor, QPainter, QPen
from PyQt5.QtWidgets import QToolTip, QWidget

_ACTION_BTN_STYLE = """
    QPushButton {
        background: #ffffff;
        border: 1px solid #d8dee8;
        border-radius: 8px;
        padding: 3px 10px;
        color: #334155;
        font-size: 11px;
        font-weight: 600;
        min-height: 26px;
    }
    QPushButton:hover {
        background: #f8fafc;
        border: 1px solid #d8dee8;
    }
    QPushButton:disabled {
        color: #94a3b8;
        background: #f1f5f9;
        border: 1px solid #d8dee8;
    }
"""

_GROUP_BOX_STYLE = """
    QGroupBox {
        border: 1px solid #d8dee8;
        border-radius: 8px;
        margin-top: 10px;
        background: #ffffff;
        color: #334155;
    }
    QGroupBox::title {
        subcontrol-origin: margin;
        left: 10px;
        padding: 0 4px;
    }
"""

_LINE_EDIT_STYLE = """
    QLineEdit {
        background: #f3f4f6;
        border: 1px solid #d8dee8;
        border-radius: 8px;
        padding: 3px 8px;
        color: #374151;
        font-size: 11px;
    }
    QLineEdit:focus {
        background: #ffffff;
        border: 1px solid #d8dee8;
    }
"""

_TABLE_STYLE = """
    QTableWidget {
        border: 1px solid #d8dee8;
        border-radius: 6px;
        gridline-color: #edf1f7;
        background: #ffffff;
        selection-background-color: #dbeafe;
        selection-color: #0f172a;
    }
    QHeaderView::section {
        background: #f8fafc;
        color: #334155;
        border: none;
        border-right: 1px solid #edf1f7;
        border-bottom: 1px solid #edf1f7;
        padding: 4px 6px;
        font-size: 11px;
        font-weight: 600;
    }
    QTableWidget::item {
        padding: 1px 6px;
    }
    QScrollBar:vertical {
        background: #f6f8fc;
        width: 8px;
        border: none;
        border-radius: 4px;
    }
    QScrollBar::handle:vertical {
        background: #cbd5e1;
        min-height: 24px;
        border-radius: 4px;
    }
    QScrollBar::add-line:vertical, QScrollBar::sub-line:vertical {
        height: 0px;
    }
    QScrollBar:horizontal {
        background: #f6f8fc;
        height: 8px;
        border: none;
        border-radius: 4px;
    }
    QScrollBar::handle:horizontal {
        background: #cbd5e1;
        min-width: 24px;
        border-radius: 4px;
    }
    QScrollBar::add-line:horizontal, QScrollBar::sub-line:horizontal {
        width: 0px;
    }
"""

_COMPACT_TAB_STYLE = """
    QTabWidget#rawDataCompactTabs::pane {
        border: 1px solid #d8dee8;
        top: -1px;
        background: #ffffff;
    }
    QTabWidget#rawDataCompactTabs QTabBar {
        left: 0px;
    }
    QTabWidget#rawDataCompactTabs QTabBar::tab {
        background: #eef1f5;
        color: #334155;
        border: 1px solid #d8dee8;
        padding: 5px 12px;
        min-width: 86px;
        font-size: 11px;
        font-weight: 600;
        margin-right: 2px;
    }
    QTabWidget#rawDataCompactTabs QTabBar::tab:selected {
        background: #ffffff;
        color: #0f172a;
        border-bottom: 1px solid #ffffff;
    }
    QTabWidget#rawDataCompactTabs QTabBar::tab:!selected {
        margin-top: 1px;
    }
    QTabWidget#rawDataCompactTabs QTabBar::tab:!selected:hover {
        background: #e2e8f0;
    }
"""

_RIGHT_TAB_STYLE = """
    QTabWidget#rightTabs::pane {
        border: 1px solid #dfe5ef;
        border-radius: 0px;
        top: -1px;
        background: #ffffff;
    }
    QTabWidget#rightTabs QTabBar {
        left: 6px;
    }
    QTabWidget#rightTabs QTabBar::tab {
        background: #eef1f5;
        color: #1f2937;
        border: 1px solid #d8dee8;
        border-radius: 0px;
        font-weight: 600;
        padding: 4px 14px;
        min-width: 56px;
        margin-right: 2px;
    }
    QTabWidget#rightTabs QTabBar::tab:selected {
        background: #ffffff;
        border-bottom: 1px solid #ffffff;
    }
    QTabWidget#rightTabs QTabBar::tab:!selected {
        margin-top: 1px;
    }
"""

_PREVIEW_NAV_BTN_STYLE = """
    QPushButton {
        background: #ffffff;
        border: 1px solid #d8dee8;
        border-radius: 6px;
        padding: 2px 12px;
        color: #334155;
        font-size: 11px;
        font-weight: 600;
        min-height: 0px;
    }
    QPushButton:hover {
        background: #f8fafc;
        border: 1px solid #d8dee8;
    }
    QPushButton:disabled {
        color: #94a3b8;
        background: #f1f5f9;
        border: 1px solid #d8dee8;
    }
"""

RESULT_HEADERS: List[str] = [
    "ID",
    "텍스트",
    "타이틀",
    "캡션",
    "이미지",
    "수식",
    "표",
    "차트",
    "각주",
    "아이템",
    "빈 텍스트 박스",
    "띄어쓰기",
    "영어",
    "기타 외국어",
    "특수문자",
    "어절",
    "박스(<2)",
    "다단 의심",
]

DEFAULT_TEXT_LABELS: List[str] = ["TEXT", "CAPTION", "TITLE", "FOOTNOTE", "ITEM"]

RESULT_COLUMN_HELP: List[str] = [
    "도서 단위 ID",
    "TEXT 라벨 박스 개수",
    "TITLE 라벨 박스 개수",
    "CAPTION 라벨 박스 개수",
    "IMAGE 라벨 박스 개수",
    "LATEX, EQUATION, FORMULA, MATH 라벨 합계",
    "TABLE 라벨 박스 개수",
    "CHART 라벨 박스 개수",
    "FOOTNOTE 라벨 박스 개수",
    "ITEM 라벨 박스 개수",
    "텍스트 검사 대상 라벨 중 비어 있는 박스 수 *텍스트 검사 대상 라벨",
    "띄어쓰기 비율 *텍스트 검사 대상 라벨",
    "영문 비율 *텍스트 검사 대상 라벨",
    "기타 외국어 비율 *텍스트 검사 대상 라벨",
    "특수문자 비율 *텍스트 검사 대상 라벨",
    "어절 수 *텍스트 검사 대상 라벨",
    "박스 수가 2개 미만인 페이지 수",
    "다단으로 의심되는 페이지 수",
]

# 분석 고정 컬럼 (라벨 카운팅 컬럼 뒤에 항상 붙는 분석 지표)
_FIXED_ANALYSIS_HEADERS: List[str] = [
    "빈 텍스트 박스",
    "띄어쓰기",
    "영어",
    "기타 외국어",
    "특수문자",
    "어절",
    "박스(<2)",
    "다단 의심",
]
_FIXED_ANALYSIS_HELP: List[str] = [
    "텍스트 검사 대상 라벨 중 비어 있는 박스 수 *텍스트 검사 대상 라벨",
    "띄어쓰기 비율 *텍스트 검사 대상 라벨",
    "영문 비율 *텍스트 검사 대상 라벨",
    "기타 외국어 비율 *텍스트 검사 대상 라벨",
    "특수문자 비율 *텍스트 검사 대상 라벨",
    "어절 수 *텍스트 검사 대상 라벨",
    "박스 수가 2개 미만인 페이지 수",
    "다단으로 의심되는 페이지 수",
]

# 동적으로 로드된 라벨 엔트리 저장 (init_dynamic_headers 호출 후 채워짐)
_label_entries: List[dict] = []


def build_result_headers(label_entries: List[dict]) -> List[str]:
    """label_entries 기반으로 결과 테이블 헤더를 생성합니다."""
    return ["ID"] + [e["title"] for e in label_entries] + _FIXED_ANALYSIS_HEADERS


def build_result_column_help(label_entries: List[dict]) -> List[str]:
    """label_entries 기반으로 결과 테이블 컬럼 설명을 생성합니다."""
    return (
        ["도서 단위 ID"]
        + [f"{e['title']} ({e['id']}) 라벨 박스 개수" for e in label_entries]
        + _FIXED_ANALYSIS_HELP
    )


def init_dynamic_headers(label_entries: List[dict]) -> None:
    """xlsx 라벨 정보를 읽어 모듈 수준 헤더 상수를 갱신합니다.

    RESULT_HEADERS, RESULT_COLUMN_HELP, DEFAULT_TEXT_LABELS 를 in-place로 교체하므로
    이 함수 호출 이후 생성되는 모든 패널은 동적 헤더를 사용합니다.
    반드시 패널 생성 전(page.py __init__ 초반)에 호출해야 합니다.
    """
    global _label_entries
    _label_entries = list(label_entries)

    new_headers = build_result_headers(label_entries)
    RESULT_HEADERS.clear()
    RESULT_HEADERS.extend(new_headers)

    new_help = build_result_column_help(label_entries)
    RESULT_COLUMN_HELP.clear()
    RESULT_COLUMN_HELP.extend(new_help)

    new_text_labels = [e["id"] for e in label_entries if e.get("is_text_analysis")]
    DEFAULT_TEXT_LABELS.clear()
    DEFAULT_TEXT_LABELS.extend(new_text_labels)
_PANEL_MARGINS = (9, 9, 9, 9)
_PANEL_SPACING = 6
_GROUP_CONTENT_MARGINS = (8, 10, 8, 8)
_GROUP_CONTENT_SPACING = 8
_FORM_PAGE_MARGINS = (8, 10, 8, 8)
_HEADER_SIDE_PADDING = 24
_HEADER_COMPACT_THRESHOLD = 84
_HEADER_LONG_MIN_WIDTH = 52
_RATIO_HEADERS = {"띄어쓰기", "영어", "기타 외국어", "특수문자"}
_ZERO_THRESHOLD_HEADERS = {"빈 텍스트 박스", "박스(<2)", "다단 의심"}
_DEFAULT_HIGHLIGHT_RULES = {
    "lower_rank_percent": 10.0,
    "upper_rank_percent": 10.0,
    "space_min": 0.20,
    "space_max": 0.33,
    "empty_text_threshold": 0.0,
    "bbox_lt_2_threshold": 0.0,
    "multicolumn_threshold": 0.0,
}
_ZERO_THRESHOLD_RULE_KEYS = {
    "빈 텍스트 박스": "empty_text_threshold",
    "박스(<2)": "bbox_lt_2_threshold",
    "다단 의심": "multicolumn_threshold",
}


def _safe_float(value: object, default: float) -> float:
    try:
        return float(str(value).strip())
    except (AttributeError, TypeError, ValueError):
        return default


def _rank_count(total_count: int, rank_percent: float) -> int:
    if total_count <= 0:
        return 0
    return max(1, ceil(total_count * (rank_percent / 100.0)))


def _format_rule_value(value: float) -> str:
    if float(value).is_integer():
        return str(int(value))
    return f"{value:.2f}".rstrip("0").rstrip(".")


def _rank_boundaries(
    values: Sequence[float],
    *,
    lower_rank_percent: float,
    upper_rank_percent: float,
) -> tuple[float, float]:
    if not values:
        return 0.0, 0.0
    sorted_values = sorted(float(value) for value in values)
    lower_rank_count = _rank_count(len(sorted_values), lower_rank_percent)
    upper_rank_count = _rank_count(len(sorted_values), upper_rank_percent)
    low_boundary = sorted_values[lower_rank_count - 1]
    high_boundary = sorted_values[-upper_rank_count]
    return low_boundary, high_boundary


def normalize_highlight_rules(rules: Optional[dict] = None) -> dict:
    merged = dict(_DEFAULT_HIGHLIGHT_RULES)
    if rules:
        merged.update(rules)
    lower_rank_percent = _safe_float(
        merged.get("lower_rank_percent", merged.get("rank_percent")),
        _DEFAULT_HIGHLIGHT_RULES["lower_rank_percent"],
    )
    upper_rank_percent = _safe_float(
        merged.get("upper_rank_percent", merged.get("rank_percent")),
        _DEFAULT_HIGHLIGHT_RULES["upper_rank_percent"],
    )
    lower_rank_percent = max(1.0, min(50.0, lower_rank_percent))
    upper_rank_percent = max(1.0, min(50.0, upper_rank_percent))
    space_min = _safe_float(merged.get("space_min"), _DEFAULT_HIGHLIGHT_RULES["space_min"])
    space_max = _safe_float(merged.get("space_max"), _DEFAULT_HIGHLIGHT_RULES["space_max"])
    space_min = max(0.0, min(1.0, space_min))
    space_max = max(0.0, min(1.0, space_max))
    if space_min > space_max:
        space_min, space_max = space_max, space_min
    return {
        "lower_rank_percent": lower_rank_percent,
        "upper_rank_percent": upper_rank_percent,
        "space_min": space_min,
        "space_max": space_max,
        "empty_text_threshold": max(
            0.0,
            _safe_float(merged.get("empty_text_threshold"), _DEFAULT_HIGHLIGHT_RULES["empty_text_threshold"]),
        ),
        "bbox_lt_2_threshold": max(
            0.0,
            _safe_float(merged.get("bbox_lt_2_threshold"), _DEFAULT_HIGHLIGHT_RULES["bbox_lt_2_threshold"]),
        ),
        "multicolumn_threshold": max(
            0.0,
            _safe_float(merged.get("multicolumn_threshold"), _DEFAULT_HIGHLIGHT_RULES["multicolumn_threshold"]),
        ),
    }


def build_result_help_tooltip() -> str:
    lines = ["결과표 컬럼 설명"]
    for header, description in zip(RESULT_HEADERS, RESULT_COLUMN_HELP):
        lines.append(f"- {header}: {description}")
    return "\n".join(lines)


def _fmt_float(value: object) -> str:
    if isinstance(value, float):
        return f"{value:.4f}"
    return str(value) if value is not None else ""


def _label_count(row: dict, label: str) -> int:
    try:
        return int(row.get(f"label_count_{label}", 0) or 0)
    except (TypeError, ValueError):
        return 0


def book_row_to_display_cells(row: dict) -> List[str]:
    if _label_entries:
        # 동적: xlsx에서 로드된 라벨 순서대로 카운트
        cells: List[str] = [str(row.get("book_name", ""))]
        for entry in _label_entries:
            cells.append(str(_label_count(row, entry["id"])))
    else:
        # 폴백: 기존 하드코딩 라벨 (xlsx 없을 때)
        formula_sum = (
            _label_count(row, "LATEX")
            + _label_count(row, "EQUATION")
            + _label_count(row, "FORMULA")
            + _label_count(row, "MATH")
        )
        cells = [
            str(row.get("book_name", "")),
            str(_label_count(row, "TEXT")),
            str(_label_count(row, "TITLE")),
            str(_label_count(row, "CAPTION")),
            str(_label_count(row, "IMAGE")),
            str(formula_sum),
            str(_label_count(row, "TABLE")),
            str(_label_count(row, "CHART")),
            str(_label_count(row, "FOOTNOTE")),
            str(_label_count(row, "ITEM")),
        ]

    cells += [
        str(int(row.get("empty_text_box_total", 0) or 0)),
        _fmt_float(row.get("space_ratio")),
        _fmt_float(row.get("english_ratio")),
        _fmt_float(row.get("other_foreign_ratio")),
        _fmt_float(row.get("special_char_ratio")),
        str(int(row.get("word_count_total", 0) or 0)),
        str(int(row.get("bbox_lt_2_pages", 0) or 0)),
        str(int(row.get("multicolumn_suspected_pages", 0) or 0)),
    ]
    return cells


def book_metric_numeric_value(header: str, row: dict) -> float:
    # 고정 분석 컬럼 (항상 동일)
    if header == "빈 텍스트 박스":
        return float(row.get("empty_text_box_total", 0) or 0)
    if header == "띄어쓰기":
        return float(row.get("space_ratio", 0.0) or 0.0)
    if header == "영어":
        return float(row.get("english_ratio", 0.0) or 0.0)
    if header == "기타 외국어":
        return float(row.get("other_foreign_ratio", 0.0) or 0.0)
    if header == "특수문자":
        return float(row.get("special_char_ratio", 0.0) or 0.0)
    if header == "어절":
        return float(row.get("word_count_total", 0) or 0)
    if header == "박스(<2)":
        return float(row.get("bbox_lt_2_pages", 0) or 0)
    if header == "다단 의심":
        return float(row.get("multicolumn_suspected_pages", 0) or 0)

    # 동적 라벨 컬럼: title → id 로 카운트 조회
    if _label_entries:
        for entry in _label_entries:
            if entry["title"] == header:
                return float(_label_count(row, entry["id"]))
        return 0.0

    # 폴백: 기존 하드코딩 (xlsx 없을 때)
    if header == "텍스트":
        return float(_label_count(row, "TEXT"))
    if header == "타이틀":
        return float(_label_count(row, "TITLE"))
    if header == "캡션":
        return float(_label_count(row, "CAPTION"))
    if header == "이미지":
        return float(_label_count(row, "IMAGE"))
    if header == "수식":
        return float(
            _label_count(row, "LATEX")
            + _label_count(row, "EQUATION")
            + _label_count(row, "FORMULA")
            + _label_count(row, "MATH")
        )
    if header == "표":
        return float(_label_count(row, "TABLE"))
    if header == "차트":
        return float(_label_count(row, "CHART"))
    if header == "각주":
        return float(_label_count(row, "FOOTNOTE"))
    if header == "아이템":
        return float(_label_count(row, "ITEM"))
    return 0.0


def format_metric_value(header: str, value: float) -> str:
    if header in _RATIO_HEADERS:
        return f"{value:.4f}"
    return str(int(value)) if float(value).is_integer() else f"{value:.2f}"


def _percentile(sorted_values: Sequence[float], ratio: float) -> float:
    if not sorted_values:
        return 0.0
    if len(sorted_values) == 1:
        return float(sorted_values[0])
    clamped = max(0.0, min(1.0, ratio))
    position = clamped * (len(sorted_values) - 1)
    lower_index = int(position)
    upper_index = min(lower_index + 1, len(sorted_values) - 1)
    weight = position - lower_index
    lower_value = float(sorted_values[lower_index])
    upper_value = float(sorted_values[upper_index])
    return lower_value + (upper_value - lower_value) * weight


def _chart_display_range(values: Sequence[float]) -> tuple[float, float]:
    if not values:
        return 0.0, 1.0
    sorted_values = sorted(float(value) for value in values)
    if len(sorted_values) == 1:
        value = sorted_values[0]
        return value - 0.5, value + 0.5
    q1 = _percentile(sorted_values, 0.25)
    q3 = _percentile(sorted_values, 0.75)
    iqr = max(q3 - q1, 0.0)
    if iqr > 0:
        lower = q1 - 1.5 * iqr
        upper = q3 + 1.5 * iqr
    else:
        p10 = _percentile(sorted_values, 0.10)
        p90 = _percentile(sorted_values, 0.90)
        lower, upper = p10, p90
    actual_min = sorted_values[0]
    actual_max = sorted_values[-1]
    lower = max(actual_min, lower)
    upper = min(actual_max, upper)
    if upper <= lower:
        lower, upper = actual_min, actual_max
    if upper <= lower:
        upper = lower + 1.0
    return lower, upper


class DistributionMiniChart(QWidget):
    def __init__(
        self,
        values: Sequence[float],
        *,
        labels: Optional[Sequence[str]] = None,
        threshold_lines: Optional[Sequence[float]] = None,
        parent: Optional[QWidget] = None,
    ) -> None:
        super().__init__(parent)
        self._values = [float(value) for value in values]
        self._labels = [str(label or "") for label in (labels or [])]
        self._threshold_lines = [float(value) for value in (threshold_lines or [])]
        self._bar_regions: List[tuple[QRect, str]] = []
        self.setMouseTracking(True)
        self.setMinimumHeight(108)
        self.setMaximumHeight(108)

    def paintEvent(self, event) -> None:  # type: ignore[override]
        del event
        painter = QPainter(self)
        painter.setRenderHint(QPainter.Antialiasing)
        self._bar_regions = []

        rect = self.rect().adjusted(8, 8, -8, -10)
        painter.fillRect(rect, QColor("#ffffff"))
        painter.setPen(QPen(QColor("#e5e7eb"), 1))
        painter.drawRoundedRect(rect, 6, 6)

        if not self._values:
            painter.setPen(QColor("#94a3b8"))
            painter.drawText(rect, Qt.AlignCenter, "데이터 없음")
            return

        display_min, display_max = _chart_display_range(self._values)
        range_width = max(display_max - display_min, 1e-9)
        q1 = _percentile(sorted(self._values), 0.25)
        q3 = _percentile(sorted(self._values), 0.75)
        iqr = max(q3 - q1, 0.0)
        if iqr > 0:
            estimated_width = max(2 * iqr / max(len(self._values) ** (1 / 3), 1.0), range_width / 8)
            bin_count = max(4, min(10, int(round(range_width / estimated_width)) or 1))
        else:
            unique_count = len({round(value, 8) for value in self._values})
            bin_count = max(4, min(10, unique_count or len(self._values)))
        bin_width = range_width / bin_count
        bins: List[List[str]] = [[] for _ in range(bin_count)]
        counts = [0] * bin_count
        value_labels = self._labels if len(self._labels) == len(self._values) else [""] * len(self._values)
        for value, book_name in zip(self._values, value_labels):
            clamped_value = min(max(value, display_min), display_max)
            index = min(int((clamped_value - display_min) / bin_width), bin_count - 1)
            counts[index] += 1
            if book_name:
                bins[index].append(book_name)

        chart_rect = rect.adjusted(10, 10, -10, -26)
        max_count = max(counts) or 1
        bar_gap = 4
        usable_width = max(chart_rect.width() - bar_gap * (bin_count - 1), bin_count)
        bar_width = max(6, usable_width // bin_count)

        painter.setPen(QPen(QColor("#cbd5e1"), 1))
        painter.drawLine(chart_rect.left(), chart_rect.bottom(), chart_rect.right(), chart_rect.bottom())
        painter.drawLine(chart_rect.left(), chart_rect.top(), chart_rect.left(), chart_rect.bottom())

        for index, count in enumerate(counts):
            bar_height = int((count / max_count) * chart_rect.height())
            x = chart_rect.left() + index * (bar_width + bar_gap)
            y = chart_rect.bottom() - bar_height
            bar_rect = QRect(x, y, bar_width, bar_height)
            painter.fillRect(bar_rect, QColor("#93c5fd"))
            start_value = display_min + index * bin_width
            end_value = display_min + (index + 1) * bin_width
            tooltip_lines = [
                f"구간: {start_value:.4f} ~ {end_value:.4f}",
                f"도서 수: {count}권",
            ]
            book_names = sorted(set(bins[index]))
            if book_names:
                tooltip_lines.append("도서명:")
                tooltip_lines.extend(book_names)
            self._bar_regions.append((bar_rect, "\n".join(tooltip_lines)))

        for threshold in self._threshold_lines:
            ratio = (threshold - display_min) / range_width
            ratio = max(0.0, min(1.0, ratio))
            x = int(chart_rect.left() + ratio * chart_rect.width())
            painter.setPen(QPen(QColor("#ef4444"), 1, Qt.DashLine))
            painter.drawLine(x, chart_rect.top(), x, chart_rect.bottom())

        painter.setPen(QColor("#64748b"))
        painter.drawText(
            rect.adjusted(10, rect.height() - 20, -10, 0),
            Qt.AlignLeft | Qt.AlignVCenter,
            f"{display_min:.2f}",
        )
        painter.drawText(
            rect.adjusted(10, rect.height() - 20, -10, 0),
            Qt.AlignCenter | Qt.AlignVCenter,
            f"{((display_min + display_max) / 2):.2f}",
        )
        painter.drawText(
            rect.adjusted(10, rect.height() - 20, -10, 0),
            Qt.AlignRight | Qt.AlignVCenter,
            f"{display_max:.2f}",
        )

    def mouseMoveEvent(self, event) -> None:  # type: ignore[override]
        for bar_rect, tooltip_text in self._bar_regions:
            if bar_rect.contains(event.pos()):
                QToolTip.showText(event.globalPos(), tooltip_text, self)
                return
        QToolTip.hideText()
        super().mouseMoveEvent(event)

    def leaveEvent(self, event) -> None:  # type: ignore[override]
        QToolTip.hideText()
        super().leaveEvent(event)


def _append_text_labels_from_rows(rows: Sequence[Sequence[object]], labels: List[str], seen: set[str]) -> None:
    id_candidates = (
        "id",
        "label id",
        "label_id",
        "라벨 id",
        "라벨id",
        "category id",
        "category_id",
        "code",
    )
    title_candidates = (
        "title",
        "label title",
        "label_title",
        "name",
        "label",
        "라벨",
        "라벨명",
        "카테고리",
        "카테고리명",
    )

    if not rows:
        return

    header = [str(cell).strip().lower() if cell is not None else "" for cell in rows[0]]

    def find_index(candidates: Sequence[str]) -> int:
        for candidate in candidates:
            if candidate in header:
                return header.index(candidate)
        return -1

    id_idx = find_index(id_candidates)
    title_idx = find_index(title_candidates)
    if id_idx < 0:
        return

    for row in rows[1:]:
        if not row:
            continue
        raw_id = row[id_idx] if id_idx < len(row) else None
        raw_title = row[title_idx] if 0 <= title_idx < len(row) else None
        label_id = str(raw_id or "").strip().upper()
        title = str(raw_title or "").strip()
        value = title or label_id
        if not label_id or label_id in seen:
            continue
        seen.add(label_id)
        labels.append(value if value == label_id else f"{value} ({label_id})")


def _xlsx_col_to_index(ref: str) -> int:
    letters = "".join(ch for ch in str(ref or "") if ch.isalpha()).upper()
    index = 0
    for char in letters:
        index = index * 26 + (ord(char) - ord("A") + 1)
    return max(0, index - 1)


def _load_text_label_options_from_xlsx_xml(workbook_path: Path) -> List[str]:
    ns = {"ss": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    rel_ns = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
    labels: List[str] = []
    seen: set[str] = set()

    with zipfile.ZipFile(workbook_path) as archive:
        workbook_root = ET.fromstring(archive.read("xl/workbook.xml"))
        rel_root = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        rel_map = {rel.attrib.get("Id", ""): rel.attrib.get("Target", "") for rel in rel_root}

        shared_strings: List[str] = []
        if "xl/sharedStrings.xml" in archive.namelist():
            shared_root = ET.fromstring(archive.read("xl/sharedStrings.xml"))
            for item in shared_root.findall("ss:si", ns):
                shared_strings.append("".join(node.text or "" for node in item.findall(".//ss:t", ns)))

        for sheet in workbook_root.findall("ss:sheets/ss:sheet", ns):
            target = rel_map.get(sheet.attrib.get(rel_ns, ""), "")
            if not target:
                continue
            sheet_path = target if target.startswith("xl/") else f"xl/{target.lstrip('/')}"
            if sheet_path not in archive.namelist():
                continue

            sheet_root = ET.fromstring(archive.read(sheet_path))
            rows: List[List[str]] = []
            for row in sheet_root.findall("ss:sheetData/ss:row", ns):
                cells: dict[int, str] = {}
                max_col = -1
                for cell in row.findall("ss:c", ns):
                    col_index = _xlsx_col_to_index(cell.attrib.get("r", ""))
                    max_col = max(max_col, col_index)
                    cell_type = cell.attrib.get("t", "")
                    value = ""
                    if cell_type == "inlineStr":
                        value = "".join(node.text or "" for node in cell.findall(".//ss:t", ns))
                    else:
                        value_node = cell.find("ss:v", ns)
                        raw_value = "" if value_node is None else str(value_node.text or "")
                        if cell_type == "s" and raw_value.isdigit():
                            shared_index = int(raw_value)
                            if 0 <= shared_index < len(shared_strings):
                                value = shared_strings[shared_index]
                        else:
                            value = raw_value
                    cells[col_index] = value
                if max_col >= 0:
                    rows.append([cells.get(index, "") for index in range(max_col + 1)])
            _append_text_labels_from_rows(rows, labels, seen)

    return labels


def load_text_label_options(project_root: Path) -> List[str]:
    label_dir = project_root / "label"
    workbook_candidates = sorted(label_dir.glob("*.xlsx"))
    if not workbook_candidates:
        return list(DEFAULT_TEXT_LABELS)

    labels: List[str] = []
    seen: set[str] = set()
    try:
        from openpyxl import load_workbook  # type: ignore
    except Exception:
        load_workbook = None

    for workbook_path in workbook_candidates:
        loaded = False
        if load_workbook is not None:
            try:
                workbook = load_workbook(workbook_path, data_only=True, read_only=True)
            except Exception:
                workbook = None
            if workbook is not None:
                try:
                    for sheet_name in workbook.sheetnames:
                        sheet = workbook[sheet_name]
                        _append_text_labels_from_rows(list(sheet.iter_rows(values_only=True)), labels, seen)
                    loaded = True
                finally:
                    workbook.close()
        if loaded:
            continue

        try:
            xml_labels = _load_text_label_options_from_xlsx_xml(workbook_path)
        except Exception:
            continue
        for value in xml_labels:
            label_id = value
            if value.endswith(")") and "(" in value:
                candidate = value.rsplit("(", 1)[-1].rstrip(")").strip()
                if candidate:
                    label_id = candidate
            label_id = label_id.upper()
            if label_id in seen:
                continue
            seen.add(label_id)
            labels.append(value)

    return labels or list(DEFAULT_TEXT_LABELS)
