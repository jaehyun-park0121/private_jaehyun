from __future__ import annotations

import logging
from datetime import date, datetime
from threading import Lock

from openpyxl import load_workbook

from ..config import AppConfig
from ..errors import MetadataError


LOGGER = logging.getLogger(__name__)


class MetadataLoader:
    def __init__(self, config: AppConfig) -> None:
        self.config = config
        self._rows_by_work_id: dict[str, dict[str, object]] | None = None
        # 병렬 실행 시 workbook 캐시 초기화 구간만 직렬화한다.
        self._load_lock = Lock()

    def ensure_loaded(self) -> None:
        if self.config.excel.workbook_path is None:
            return
        if self._rows_by_work_id is not None:
            return

        with self._load_lock:
            if self._rows_by_work_id is not None:
                return
            try:
                self._rows_by_work_id = self._load_rows()
            except Exception as exc:
                LOGGER.warning("메타데이터 로딩 중 오류가 발생하여 빈 메타데이터로 진행합니다: %s", exc)
                self._rows_by_work_id = {}

    def get_document_metadata(self, work_id: str) -> dict:
        if self.config.excel.workbook_path is None:
            LOGGER.warning("메타데이터 엑셀 파일이 설정되지 않아 빈 메타데이터로 진행합니다.")
            return _empty_metadata(work_id)

        self.ensure_loaded()
        row = (self._rows_by_work_id or {}).get(work_id)
        if row is None:
            LOGGER.warning("work_id '%s'에 해당하는 메타데이터 행이 없어 빈 메타데이터로 진행합니다.", work_id)
            return _empty_metadata(work_id)

        columns = self.config.excel.columns
        return {
            "work_id": _normalize_text(row.get(columns["work_id"], "")),
            "title": _normalize_text(row.get(columns["title"], "")),
            "identifiers": {
                "isbn": _normalize_text(row.get(columns["isbn"], "")),
            },
            "source_format": _normalize_text(row.get(columns["source_format"], "")),
            "category_1": _normalize_text(row.get(columns["category_1"], "")),
            "category_2": _normalize_text(row.get(columns["category_2"], "")),
            "category_3": _normalize_text(row.get(columns["category_3"], "")),
            "author": _normalize_text(row.get(columns["author"], "")),
            "publisher": _normalize_text(row.get(columns["publisher"], "")),
            "published_date": _normalize_date(row.get(columns["published_date"], "")),
        }

    def _load_rows(self) -> dict[str, dict[str, object]]:
        if self.config.excel.workbook_path is None:
            return {}

        workbook = load_workbook(self.config.excel.workbook_path, read_only=True, data_only=True)
        try:
            if self.config.excel.sheet_name not in workbook.sheetnames:
                raise MetadataError(f"엑셀 시트 '{self.config.excel.sheet_name}'를 찾을 수 없습니다.")

            sheet = workbook[self.config.excel.sheet_name]
            headers = [
                cell.value if cell.value is not None else ""
                for cell in sheet[self.config.excel.header_row]
            ]

            id_header = self.config.excel.columns["work_id"]
            if id_header not in headers:
                LOGGER.warning(
                    "메타데이터 ID 헤더 '%s'가 없어 빈 메타데이터로 진행합니다.",
                    id_header,
                )
                return {}

            missing_headers = [
                header_name
                for field_name, header_name in self.config.excel.columns.items()
                if field_name != "work_id" and header_name not in headers
            ]
            if missing_headers:
                LOGGER.warning(
                    "메타데이터 헤더가 없어 해당 값은 빈 문자열로 처리합니다: %s",
                    ", ".join(missing_headers),
                )

            rows_by_work_id: dict[str, dict[str, object]] = {}
            data_start_row = self.config.excel.header_row + 1
            for row_values in sheet.iter_rows(min_row=data_start_row, values_only=True):
                row_dict = {header: value for header, value in zip(headers, row_values)}
                work_id = str(row_dict.get(id_header, "") or "").strip()
                if not work_id:
                    continue
                if work_id in rows_by_work_id:
                    LOGGER.warning("메타데이터 ID '%s'가 중복되어 마지막 행을 사용합니다.", work_id)
                rows_by_work_id[work_id] = row_dict
            return rows_by_work_id
        finally:
            workbook.close()


def _normalize_date(value: object) -> str:
    if value in (None, ""):
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def _normalize_text(value: object) -> str:
    if value == "":
        return ""
    return str(value).strip()


def _empty_metadata(work_id: str) -> dict:
    return {
        "work_id": work_id,
        "title": "",
        "identifiers": {
            "isbn": "",
        },
        "source_format": "",
        "category_1": "",
        "category_2": "",
        "category_3": "",
        "author": "",
        "publisher": "",
        "published_date": "",
    }
